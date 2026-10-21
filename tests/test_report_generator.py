"""Unit tests for Excel report export."""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from src.configuration import AppConfig, OUTPUT_DETAIL_COLUMNS
from src.productivity_calculator import calculate_productivity
from src.report_generator import generate_report_bytes
from src.validators import build_standard_lookup

SAMPLES = Path(__file__).resolve().parents[1] / "samples"

PROD_MAP = {
    "task_type": "TaskType",
    "change_or_new": "ChangeOrNew",
    "complexity": "Complexity",
    "effort_manday": "EffortManday",
    "assignee": "Assignee",
    "finished_date": "Finished Date",
}

STD_MAP = {
    "work_category": "Work Category",
    "new_change": "New/Change",
    "difficulty": "Difficulty",
    "standard_effort": "Standard Development Man-Day",
}


def _sample_result():
    prod_df = pd.DataFrame(
        [
            {
                "_source_row": 2,
                "Key": "T-1",
                "TaskType": "TaskType_MI",
                "ChangeOrNew": "Change",
                "Complexity": "Complexity_C1",
                "EffortManday": "EffortManday_1",
                "Assignee": "LE DUC ANH Professional/(Part) Sales & CS",
                "Finished Date": "2026-10-15",
            },
            {
                "_source_row": 3,
                "Key": "T-2",
                "TaskType": "TaskType_ZZ",
                "ChangeOrNew": "Change",
                "Complexity": "Complexity_C1",
                "EffortManday": "EffortManday_X",
                "Assignee": "SOMEONE Intern/Dept",
                "Finished Date": "2026-10-16",
            },
        ]
    )
    std_df = pd.DataFrame(
        [
            {
                "_source_row": 2,
                "Work Category": "MI",
                "New/Change": "Change",
                "Difficulty": "C1",
                "Standard Development Man-Day": 2,
            }
        ]
    )
    lookup = build_standard_lookup(std_df, STD_MAP)
    result = calculate_productivity(prod_df, PROD_MAP, lookup, AppConfig())
    return result, lookup


def test_export_worksheets_and_columns():
    result, lookup = _sample_result()
    data = generate_report_bytes(result, lookup, AppConfig())
    wb = openpyxl.load_workbook(io.BytesIO(data))

    assert "Productivity Detail" in wb.sheetnames
    assert "Summary" in wb.sheetnames
    assert "Validation Issues" in wb.sheetnames
    assert "Standard Reference" in wb.sheetnames

    ws = wb["Productivity Detail"]
    headers = [c.value for c in ws[1]]
    for col in OUTPUT_DETAIL_COLUMNS:
        assert col in headers

    # Excluded from export
    for col in ("Source Row", "Employee Coefficient", "Validation Status"):
        assert col not in headers

    # Column order: required detail columns appear in the specified order.
    positions = [headers.index(c) for c in OUTPUT_DETAIL_COLUMNS]
    assert positions == sorted(positions)

    # Numeric cells stored as numbers
    # Row 2 = first task (valid)
    emp_col = headers.index("Employee Name") + 1
    effort_col = headers.index("Total Effort (MD)") + 1
    prod_col = headers.index("Productivity") + 1
    specialist_col = headers.index("Effort For Specialist") + 1

    assert ws.cell(2, emp_col).value == "LE DUC ANH"
    assert isinstance(ws.cell(2, effort_col).value, (int, float))
    assert ws.cell(2, effort_col).value == pytest.approx(1.0)
    assert ws.cell(2, prod_col).value == "200.00%"
    assert isinstance(ws.cell(2, specialist_col).value, (int, float))
    assert ws.cell(2, specialist_col).value == pytest.approx(1.43)

    # Invalid task preserved
    assert ws.max_row >= 3
    issues_col = headers.index("Validation Issues") + 1
    assert ws.cell(3, issues_col).value  # has issues text

    # Reopen successfully
    wb2 = openpyxl.load_workbook(io.BytesIO(data))
    assert wb2.sheetnames


def test_export_with_sample_template():
    template_path = SAMPLES / "Output_Template.xlsx"
    if not template_path.exists():
        pytest.skip("Sample Output_Template.xlsx not available")

    result, lookup = _sample_result()
    data = generate_report_bytes(
        result,
        lookup,
        AppConfig(),
        template_bytes=template_path.read_bytes(),
        template_sheet="Sheet1",
    )
    wb = openpyxl.load_workbook(io.BytesIO(data))
    # Primary sheet renamed or kept
    assert any(
        name in wb.sheetnames for name in ("Productivity Detail", "Sheet1")
    )
    # Support sheets still present
    assert "Summary" in wb.sheetnames

    # Find detail sheet
    detail_name = "Productivity Detail" if "Productivity Detail" in wb.sheetnames else "Sheet1"
    ws = wb[detail_name]
    headers = [c.value for c in ws[1]]
    assert "Employee Name" in headers
    assert "Productivity" in headers
    assert "Effort For Specialist" in headers

    # Values written
    name_col = headers.index("Employee Name") + 1
    assert ws.cell(2, name_col).value == "LE DUC ANH"
    prod_col = headers.index("Productivity") + 1
    assert str(ws.cell(2, prod_col).value).endswith("%")


def test_export_from_real_sample_files():
    """End-to-end smoke test against repository sample workbooks."""
    prod_path = SAMPLES / "Productivity Report.xlsx"
    std_path = SAMPLES / "Standard.xlsx"
    if not prod_path.exists() or not std_path.exists():
        pytest.skip("Sample workbooks not available")

    from src.excel_reader import read_table

    prod_df = read_table(prod_path)
    std_df = read_table(std_path)
    lookup = build_standard_lookup(std_df, STD_MAP)
    result = calculate_productivity(prod_df, PROD_MAP, lookup, AppConfig())

    assert result.summary.total_tasks == 9
    assert result.summary.valid_productivity_tasks > 0

    data = generate_report_bytes(result, lookup, AppConfig())
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "Productivity Detail" in wb.sheetnames
    assert wb["Productivity Detail"].max_row == result.summary.total_tasks + 1
