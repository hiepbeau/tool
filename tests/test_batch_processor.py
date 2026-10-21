"""Tests for multi-file batch processing."""

from __future__ import annotations

import io

import openpyxl
import pandas as pd
import pytest

from src.batch_processor import (
    FileInput,
    calculate_batch,
    per_file_summary_table,
    tag_source_file,
)
from src.configuration import AppConfig
from src.productivity_calculator import calculate_productivity
from src.report_generator import generate_consolidated_report_bytes, generate_report_bytes
from src.validators import build_standard_lookup

PROD_MAP = {
    "task_type": "TaskType",
    "change_or_new": "ChangeOrNew",
    "complexity": "Complexity",
    "effort_manday": "EffortManday",
    "assignee": "Assignee",
    "finished_date": "Finished Date",
}

STD_MAP = {
    "work_category": "Work Category",
    "new_change": "New/Change",
    "difficulty": "Difficulty",
    "standard_effort": "Standard Development Man-Day",
}


def _task(key, effort="EffortManday_1", complexity="Complexity_C1", **extra):
    row = {
        "_source_row": 2,
        "Key": key,
        "TaskType": "TaskType_MI",
        "ChangeOrNew": "Change",
        "Complexity": complexity,
        "EffortManday": effort,
        "Assignee": "LE DUC ANH Professional/(Part) Sales & CS",
        "Finished Date": "2026-10-15",
    }
    row.update(extra)
    return row


def _standard():
    return pd.DataFrame(
        [
            {
                "_source_row": 2,
                "Work Category": "MI",
                "New/Change": "Change",
                "Difficulty": "C1",
                "Standard Development Man-Day": 2,
            },
            {
                "_source_row": 3,
                "Work Category": "MI",
                "New/Change": "Change",
                "Difficulty": "C2",
                "Standard Development Man-Day": 3,
            },
        ]
    )


def test_tag_source_file():
    lookup = build_standard_lookup(_standard(), STD_MAP)
    result = calculate_productivity(pd.DataFrame([_task("A")]), PROD_MAP, lookup)
    tagged = tag_source_file(result, "Report_A.xlsx")
    assert "Source File" in tagged.detail.columns
    assert tagged.detail.iloc[0]["Source File"] == "Report_A.xlsx"
    assert tagged.detail.iloc[0]["Productivity"] == pytest.approx(200.0)


def test_calculate_batch_independent_and_combined():
    lookup = build_standard_lookup(_standard(), STD_MAP)
    files = [
        FileInput(
            name="A.xlsx",
            dataframe=pd.DataFrame([_task("A1")]),
            sheet_name="Sheet1",
        ),
        FileInput(
            name="B.xlsx",
            dataframe=pd.DataFrame(
                [_task("B1", effort="EffortManday_2", complexity="Complexity_C2", _source_row=2)]
            ),
            sheet_name="Sheet1",
        ),
    ]
    batch = calculate_batch(files, PROD_MAP, lookup, AppConfig())
    assert batch.file_count == 2
    assert batch.errors == []
    assert batch.combined is not None

    # Independent calculations
    assert batch.file_results[0].result.detail.iloc[0]["Productivity"] == pytest.approx(200.0)
    assert batch.file_results[1].result.detail.iloc[0]["Productivity"] == pytest.approx(150.0)

    # Combined preserves all tasks and Source File
    assert len(batch.combined.detail) == 2
    assert set(batch.combined.detail["Source File"]) == {"A.xlsx", "B.xlsx"}

    # Combined weighted: (2+3)/(1+2)*100 = 166.666...
    assert batch.combined.summary.total_tasks == 2
    assert batch.combined.summary.valid_productivity_tasks == 2
    assert batch.combined.summary.weighted_productivity == pytest.approx(500 / 3)
    assert batch.combined.summary.total_actual_effort == pytest.approx(3.0)
    assert batch.combined.summary.total_standard_effort == pytest.approx(5.0)


def test_combine_does_not_average_file_percentages():
    lookup = build_standard_lookup(_standard(), STD_MAP)
    files = [
        FileInput("A.xlsx", pd.DataFrame([_task("A")])),
        FileInput(
            "B.xlsx",
            pd.DataFrame([_task("B", effort="EffortManday_2", complexity="Complexity_C2")]),
        ),
    ]
    batch = calculate_batch(files, PROD_MAP, lookup)
    # File averages would be (200+150)/2 = 175; combined weighted is 166.66
    assert batch.combined.summary.weighted_productivity != pytest.approx(175.0)


def test_per_file_summary_table():
    lookup = build_standard_lookup(_standard(), STD_MAP)
    batch = calculate_batch(
        [
            FileInput("A.xlsx", pd.DataFrame([_task("A")]), "Sheet1"),
            FileInput("B.xlsx", pd.DataFrame([_task("B", TaskType="TaskType_ZZ")]), "Sheet1"),
        ],
        PROD_MAP,
        lookup,
    )
    table = per_file_summary_table(batch, AppConfig())
    assert list(table["Source File"]) == ["A.xlsx", "B.xlsx"]
    assert int(table.loc[0, "Valid Tasks"]) == 1
    assert int(table.loc[1, "Missing Standard"]) == 1


def test_consolidated_export_includes_source_file_and_sheets():
    lookup = build_standard_lookup(_standard(), STD_MAP)
    batch = calculate_batch(
        [
            FileInput("A.xlsx", pd.DataFrame([_task("A")])),
            FileInput("B.xlsx", pd.DataFrame([_task("B", TaskType="TaskType_ZZ")])),
        ],
        PROD_MAP,
        lookup,
    )
    data = generate_consolidated_report_bytes(batch, AppConfig())
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "Productivity Detail" in wb.sheetnames
    assert "Summary" in wb.sheetnames
    assert "Per-File Summary" in wb.sheetnames
    assert "Validation Issues" in wb.sheetnames

    headers = [c.value for c in wb["Productivity Detail"][1]]
    assert "Source File" in headers
    assert "Employee Name" in headers

    # Individual export still works
    individual = generate_report_bytes(batch.file_results[0].result, lookup, AppConfig())
    wb2 = openpyxl.load_workbook(io.BytesIO(individual))
    assert "Productivity Detail" in wb2.sheetnames


def test_single_file_batch_still_works():
    """Backward compatible: one file behaves like before."""
    lookup = build_standard_lookup(_standard(), STD_MAP)
    batch = calculate_batch(
        [FileInput("Only.xlsx", pd.DataFrame([_task("T1")]))],
        PROD_MAP,
        lookup,
    )
    assert batch.file_count == 1
    assert len(batch.combined.detail) == 1
    assert batch.combined.detail.iloc[0]["Productivity"] == pytest.approx(200.0)
    assert batch.combined.detail.iloc[0]["Effort For Specialist"] == pytest.approx(1.43)
